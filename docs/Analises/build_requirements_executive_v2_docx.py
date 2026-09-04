"""Gera relatório executivo V2 de requisitos do Canal de Denúncias."""
from collections import Counter
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from build_requirements_executive_docx import (
    GOLD,
    GRAY_D,
    GREEN,
    NAVY,
    NAVY_HEX,
    RED,
    TEAL,
    add_bullet,
    add_heading,
    add_numbered,
    add_paragraph,
    add_table,
    set_cell_shading,
)


ROOT = Path(__file__).resolve().parent.parent
PROJECT = ROOT / "cidadania-canal-denuncias"
OUT = Path(__file__).parent / "RELATORIO_EXECUTIVO_REQUISITOS_DENUNCIAS_V2.docx"
REPORT_DATE = date(2026, 7, 24)


RF = [
    ("RF-01", "Apresentar acolhimento e permitir iniciar uma denúncia.", "Implementado",
     "src/app/components/acolhimento", "A ação “Denunciar” inicia o wizard."),
    ("RF-02", "Apresentar vídeo explicativo de 45 segundos.", "Somente visual",
     "acolhimento.html#video-institucional", "Há placeholder; não existe mídia nem ação de reprodução."),
    ("RF-03", "Oferecer fluxos alternativos para denúncia de órgão público e contato com a Ouvidoria.",
     "Somente visual", "acolhimento.html", "Os cards não possuem evento, rota ou link."),
    ("RF-04", "Conduzir o usuário em wizard com seis etapas, revisão e confirmação.", "Implementado",
     "app.html; ComplaintService", "Estado controlado por signals, sem rotas por etapa."),
    ("RF-05", "Preservar os dados ao avançar, voltar e editar etapas durante a sessão.", "Implementado",
     "ComplaintService; componentes OnDestroy", "Não persiste após recarregar/fechar o navegador."),
    ("RF-06", "Selecionar uma ou mais categorias de irregularidade entre seis opções.", "Implementado",
     "complaint.model.ts#IRREGULARIDADES", "Jornada, CTPS, EPI, assédio, salários e outros."),
    ("RF-07", "Registrar relato textual livre.", "Implementado",
     "step-irregularidades", "Relato ou ao menos uma categoria é obrigatório no backend."),
    ("RF-08", "Gravar relato em áudio e, se aplicável, produzir transcrição.", "Parcial",
     "step-irregularidades; complaint.model.ts", "O Blob é criado, mas não é anexado, transmitido, transcrito ou revisado."),
    ("RF-09", "Registrar período, modalidade, número de prejudicados, funções/setores e nomes/dados.",
     "Implementado", "step-ocorrencias", "Todos os campos são opcionais na lógica atual."),
    ("RF-10", "Indicar grupos vulneráveis envolvidos.", "Implementado",
     "step-evidencias", "Opções atuais: crianças/adolescentes e pessoas com deficiência."),
    ("RF-11", "Anexar múltiplas imagens, áudios, vídeos e documentos.", "Implementado",
     "step-evidencias; middleware/upload.js", "O backend limita quantidade, tamanho e MIME."),
    ("RF-12", "Listar e remover anexos antes do envio, com prévia quando aplicável.", "Implementado",
     "step-evidencias", "Prévia em base64 é criada apenas para imagens."),
    ("RF-13", "Informar se outro órgão já foi procurado.", "Implementado",
     "step-evidencias", "Resposta Sim/Não é opcional."),
    ("RF-14", "Permitir denúncia anônima ou identificada.", "Implementado",
     "step-identificacao", "O modo inicial é anônimo."),
    ("RF-15", "Coletar nome, e-mail e telefone quando a pessoa optar por identificação.", "Parcial",
     "step-identificacao; complaint.service.js", "Campos aparecem, mas não são obrigatórios nem validados."),
    ("RF-16", "Coletar UF e município do fato.", "Implementado",
     "step-local; server/services/complaint.service.js", "São os únicos dados de localização obrigatórios no backend."),
    ("RF-17", "Coletar nome, endereço e CNPJ da empresa denunciada.", "Parcial",
     "step-local", "CNPJ recebe máscara, mas não há validação dos dígitos verificadores."),
    ("RF-18", "Exibir revisão da denúncia e permitir editar cada etapa.", "Parcial",
     "step-revisao", "A revisão omite vários campos, dados de contato, áudio, CNPJ e endereço."),
    ("RF-19", "Serializar denúncia e anexos em multipart e enviar ao BFF.", "Implementado",
     "complaint-api.client.ts", "POST /api/denuncias."),
    ("RF-20", "Validar JSON e conteúdo mínimo antes do encaminhamento.", "Implementado",
     "complaint.controller.js; complaint.service.js", "Exige UF, município e categoria ou relato."),
    ("RF-21", "Restringir anexos por quantidade, tamanho e tipos permitidos.", "Implementado",
     "middleware/upload.js", "Padrão: 10 arquivos, 20 MB por arquivo, 15 MIME types."),
    ("RF-22", "Verificar anexos em scanner antimalware antes do encaminhamento.", "Implementado",
     "clamav.client.js; complaint.controller.js", "Falha fechada; infraestrutura real ainda precisa ser provisionada."),
    ("RF-23", "Encaminhar a denúncia à API interna do MPT sem expor token no navegador.", "Implementado",
     "mpt-api.client.js", "Token opcional fica no BFF; timeout de 30 segundos."),
    ("RF-24", "Gerar protocolo MPT-XXXXXXXX e confirmar somente após aceite da API interna.", "Implementado",
     "complaint.service.js; complaint.controller.js", "Alfabeto reduzido; não confirma falso sucesso."),
    ("RF-25", "Exibir confirmação e protocolo retornado pelo backend.", "Implementado",
     "confirmation; complaint-api.client.ts", "Resposta 201 sem protocolo é tratada como falha."),
    ("RF-26", "Exibir erro de envio, manter a revisão e permitir nova tentativa.", "Implementado",
     "ComplaintService; step-revisao", "Mensagem atual é genérica para todas as falhas."),
    ("RF-27", "Permitir iniciar uma nova denúncia após confirmação.", "Implementado",
     "confirmation; ComplaintService.reset", "Reinicia o estado em memória."),
    ("RF-28", "Disponibilizar informações do endpoint, health check e documentação da API.",
     "Implementado", "complaint.routes.js; index.js", "GET /api/denuncias/info, /health e /api-docs."),
    ("RF-29", "Permitir acompanhamento posterior por protocolo.", "Ausente",
     "Textos de acolhimento/confirmacao", "O protocolo é exibido, mas não existe consulta ou acompanhamento."),
    ("RF-30", "Executar uma ação de acessibilidade pelo botão do cabeçalho.", "Somente visual",
     "app.html#btn-accessibility", "Botão possui rótulo, mas não possui ação."),
]


RNF = [
    ("RNF-01", "Arquitetura separada em SPA Angular e BFF Express.", "Atendido",
     "src/; server/", "Integrações e token permanecem no BFF."),
    ("RNF-02", "Ambiente reproduzível com Node 22.22.3 e npm 11.11.0 via Volta.", "Atendido",
     "package.json", "Frontend e backend possuem lockfiles independentes."),
    ("RNF-03", "Interface mobile-first e responsiva.", "Atendido",
     "spec_design.md; styles.css", "Há projeto Playwright Pixel 7 e desktop."),
    ("RNF-04", "Acessibilidade WCAG 2.1 A/AA sem violações sérias/críticas automatizadas.", "Parcial",
     "axe-core; e2e", "Automação aprovada; avaliação manual e eMAG não estão evidenciados."),
    ("RNF-05", "Compatibilidade com Chromium, Firefox e Safari.", "Parcial",
     "playwright.config.ts", "WebKit é coberto; Safari real em macOS/iPhone não foi homologado."),
    ("RNF-06", "Bundle inicial abaixo de 500 kB bruto e 120 kB transferido.", "Atendido",
     "angular.json; relatório de testes", "Última evidência: 384,38 kB / 85,26 kB."),
    ("RNF-07", "CORS restrito e cabeçalhos de segurança.", "Atendido",
     "server/index.js", "Helmet ativo e FRONTEND_URL exclusiva em produção."),
    ("RNF-08", "Segredos e detalhes da API interna restritos ao backend.", "Atendido",
     "mpt-api.client.js; env.template", "MPT_API_TOKEN não é enviado ao frontend."),
    ("RNF-09", "Limite de submissões consistente entre três réplicas.", "Parcial",
     "express-rate-limit; Redis store", "Código preparado; Redis real ainda não homologado."),
    ("RNF-10", "Antimalware obrigatório e fail closed.", "Parcial",
     "clamav.client.js", "Código e simulador aprovados; ClamAV real ainda não homologado."),
    ("RNF-11", "Anexos fora da memória, com nomes aleatórios e limpeza após requisição.", "Atendido",
     "middleware/upload.js", "Não há rotina de limpeza de órfãos após queda abrupta do processo."),
    ("RNF-12", "Limites de upload configuráveis e seguros.", "Atendido",
     "MAX_FILE_SIZE; MAX_FILES", "Padrão: 10 × 20 MB; limite total não está formalizado."),
    ("RNF-13", "Validar assinatura real do arquivo, não apenas MIME informado.", "Não atendido",
     "middleware/upload.js", "O filtro confia no Content-Type fornecido pelo cliente."),
    ("RNF-14", "Separar liveness e readiness das dependências externas.", "Não atendido",
     "GET /health", "Health sempre responde ok; não verifica Redis nem ClamAV."),
    ("RNF-15", "Nunca apresentar sucesso quando a entrega não for confirmada.", "Atendido",
     "controller; ComplaintApiClient", "201 só após aceite e protocolo válido."),
    ("RNF-16", "Não registrar conteúdo sensível da denúncia nos logs.", "Parcial",
     "controller; morgan", "Payload não é logado; formato combined registra IP e metadados HTTP."),
    ("RNF-17", "Reter logs por 10 dias.", "Não atendido no código",
     "Requisito operacional informado", "Depende da plataforma de logs; não há configuração de implantação."),
    ("RNF-18", "Disponibilidade de 99,5%.", "Não formalizado",
     "Requisito operacional informado", "Sem SLI, janela, orçamento de erro ou alertas definidos."),
    ("RNF-19", "Suportar 100 denúncias/h, 500 usuários, 100 uploads simultâneos e 3 réplicas.",
     "Não comprovado", "Requisito operacional informado", "Não há teste de carga com Redis/ClamAV/API reais."),
    ("RNF-20", "Aplicar timeouts e cancelamento em integrações.", "Parcial",
     "MPT 30s; ClamAV 30s", "Fetch frontend não possui timeout/AbortController."),
    ("RNF-21", "Observabilidade com correlação, métricas e alertas.", "Não atendido",
     "server/index.js", "Há logs textuais, mas não correlation ID, métricas ou tracing."),
    ("RNF-22", "Evitar duplicidade em retentativas.", "Não atendido",
     "Fluxo de envio", "Não há Idempotency-Key nem regra acordada com a API MPT."),
    ("RNF-23", "Qualidade automatizada e regressão multi-browser.", "Atendido",
     "Vitest; Supertest; Playwright; ESLint", "Evidência: 30 backend, 8 frontend e 16 E2E aprovados."),
    ("RNF-24", "Dependências de produção sem vulnerabilidades conhecidas na auditoria.", "Atendido",
     "npm audit --omit=dev", "Evidência local de 24/07/2026: zero em frontend e backend."),
    ("RNF-25", "Pipeline de entrega com gates completos.", "Parcial",
     ".github/workflows/npm-audit.yml", "Há gate de auditoria; pipeline integral não está documentado."),
    ("RNF-26", "Privacidade/LGPD com base legal, transparência, retenção e direitos definidos.",
     "Não formalizado", "UI; documentação", "Há promessas de sigilo, mas não política/termos/base legal no codebase."),
    ("RNF-27", "Implantação reproduzível e reversível.", "Não atendido",
     "Repositório", "Não existem Dockerfile, IaC, manifests ou definição da plataforma."),
]


BUSINESS_RULES = [
    ("RN-01", "A denúncia inicia como anônima.", "ComplaintService.initialComplaint"),
    ("RN-02", "UF e município são obrigatórios no envio.", "validateComplaint"),
    ("RN-03", "É obrigatório informar ao menos uma irregularidade ou relato textual.", "validateComplaint"),
    ("RN-04", "Demais campos do wizard são opcionais na lógica atual.", "Frontend/backend"),
    ("RN-05", "Cada anexo pode ter até 20 MB e cada denúncia até 10 anexos.", "middleware/upload.js"),
    ("RN-06", "Somente os MIME types da allowlist backend são aceitos.", "middleware/upload.js"),
    ("RN-07", "Qualquer anexo deve ser aprovado pelo ClamAV; indisponibilidade bloqueia o envio.", "controller"),
    ("RN-08", "O protocolo é gerado pelo BFF e só é devolvido após aceite da API MPT.", "controller"),
    ("RN-09", "Denúncia anônima produz sigilo=true; identificada produz sigilo=false.", "buildComplaintPayload"),
    ("RN-10", "O rate limit padrão é 20 tentativas por IP a cada 15 minutos.", "server/index.js"),
    ("RN-11", "Swagger fica desabilitado em produção salvo opt-in.", "ENABLE_API_DOCS"),
    ("RN-12", "Produção exige MPT_API_URL, FRONTEND_URL, REDIS_URL e CLAMAV_HOST.", "assertProductionConfiguration"),
    ("RN-13", "Falha da API MPT retorna 502; configuração/antimalware indisponível retorna 503.", "controller"),
    ("RN-14", "Arquivos temporários são excluídos ao fim da requisição.", "cleanupUploadedFiles"),
]


DATA_FIELDS = [
    ("irregularidades", "Lista", "Condicional", "Uma ou mais categorias; alternativa ao relato.", "Enviado"),
    ("relato_texto", "Texto", "Condicional", "Texto livre; alternativa às categorias.", "Enviado"),
    ("relato_audio", "Blob", "Opcional/indefinido", "Gravado no navegador.", "Não transmitido corretamente"),
    ("relato_audio_transcricao", "Texto", "Indefinido", "Existe no modelo, sem geração.", "Não produzido"),
    ("periodo_ocorrencia", "Texto", "Opcional", "Período informado livremente.", "Enviado"),
    ("modalidade_trabalho", "Texto", "Opcional", "Presencial/remoto etc.", "Enviado"),
    ("numero_prejudicados", "Texto", "Opcional", "Não é numérico estruturado.", "Enviado"),
    ("funcoes_setores", "Texto", "Opcional", "Funções/setores afetados.", "Enviado; omitido na revisão"),
    ("nomes_dados", "Texto sensível", "Opcional", "Nomes, e-mails e dados de terceiros.", "Enviado; omitido na revisão"),
    ("grupos_vulneraveis", "Lista", "Opcional", "Duas categorias atuais.", "Enviado"),
    ("anexos", "Arquivo", "Opcional", "Até 10 × 20 MB.", "Multipart separado"),
    ("procurou_outro_orgao", "Boolean/nulo", "Opcional", "Sim, não ou não informado.", "Enviado"),
    ("tipo_identificacao", "Enum", "Padrão anônimo", "anonimo | identificado.", "Enviado"),
    ("nome_completo", "Texto pessoal", "Não validado", "Exibido quando identificado.", "Enviado mesmo se voltar a anônimo"),
    ("email", "Texto pessoal", "Não validado", "Input type=email sem bloqueio de avanço.", "Enviado"),
    ("telefone", "Texto pessoal", "Não validado", "Sem máscara ou validação.", "Enviado"),
    ("uf", "Enum", "Obrigatório", "27 UFs.", "Validado apenas como não vazio"),
    ("municipio", "Texto", "Obrigatório", "Sem lista oficial/código IBGE.", "Validado apenas como não vazio"),
    ("nome_empresa", "Texto", "Opcional", "Nome informado livremente.", "Enviado"),
    ("endereco_empresa", "Texto", "Opcional", "Endereço/ponto de referência.", "Enviado; omitido na revisão"),
    ("cnpj_empresa", "Texto", "Opcional", "Máscara 00.000.000/0000-00.", "Sem dígito verificador; omitido na revisão"),
]


CONFLICTS = [
    ("C-01", "Crítico", "Áudio aparece como gravado com sucesso, mas não é transmitido nem transcrito.",
     "Decidir se áudio pertence ao escopo; implementar ponta a ponta ou retirar a promessa."),
    ("C-02", "Crítico", "A mensagem “Sua voz não será identificada” não possui mecanismo técnico de anonimização de voz.",
     "Validar texto com negócio/jurídico e definir tratamento técnico."),
    ("C-03", "Crítico", "Ao trocar de identificado para anônimo, nome/e-mail/telefone permanecem no estado e são enviados.",
     "Limpar dados pessoais ao selecionar anonimato ou obter regra formal diferente."),
    ("C-04", "Crítico", "A UI afirma sigilo dos dados pessoais, mas o payload define sigilo=false para identificado.",
     "Definir o significado jurídico/técnico de sigilo e anonimato."),
    ("C-05", "Alto", "A UI marca UF/município como obrigatórios, mas permite chegar à revisão e tentar enviar vazios.",
     "Aplicar validação antes da navegação e mensagem por campo."),
    ("C-06", "Alto", "A documentação prevê validação de CNPJ; o código faz somente máscara.",
     "Definir se CNPJ é apenas informativo ou deve ter dígito verificador."),
    ("C-07", "Alto", "A revisão promete conferir os dados, mas omite contato, endereço, CNPJ, nomes/dados, funções e áudio.",
     "Definir o conjunto obrigatório da revisão e exibir dados sensíveis com cuidado."),
    ("C-08", "Alto", "O seletor frontend aceita image/*, audio/* e video/*; o backend aceita uma lista menor.",
     "Alinhar accept, mensagens e allowlist backend."),
    ("C-09", "Alto", "O protocolo é anunciado para acompanhamento, mas não existe consulta de protocolo.",
     "Retirar promessa ou definir serviço de acompanhamento."),
    ("C-10", "Médio", "Vídeo explicativo, Ouvidoria e denúncia de órgão público são elementos visuais sem ação.",
     "Implementar destinos ou remover/desabilitar com explicação."),
    ("C-11", "Médio", "O botão de acessibilidade não executa ação.", "Definir função ou remover o controle."),
    ("C-12", "Médio", "README raiz indica porta 4200; scripts usam 4201.", "Atualizar documentação."),
    ("C-13", "Médio", "planejamento/spec_design citam Angular 21; package usa Angular 22.", "Atualizar baseline documental."),
    ("C-14", "Médio", "AGENTS.md diz memoryStorage; código usa diskStorage.", "Atualizar instruções do repositório."),
    ("C-15", "Médio", "AGENTS.md ainda diz que o frontend simula sucesso; o comportamento foi corrigido.",
     "Atualizar pitfalls para evitar decisões erradas."),
    ("C-16", "Médio", "server/README orienta copiar .env.example; o arquivo existente é env.template.",
     "Padronizar nome e instrução."),
    ("C-17", "Médio", "Swagger descreve 503 só por configuração MPT; 503 também ocorre por ClamAV.",
     "Atualizar contrato OpenAPI."),
    ("C-18", "Médio", "A navegação do stepper permite pular etapas sem critério de conclusão.",
     "Definir liberdade de navegação versus bloqueio por validação."),
    ("C-19", "Médio", "A promessa de anonimato total convive com logs HTTP por IP e rate limit por IP.",
     "Definir anonimato perante o processo, infraestrutura e política de logs."),
]


GAPS = [
    ("L-01", "P0", "Validar assinatura/magic bytes dos arquivos.", "Segurança de upload"),
    ("L-02", "P0", "Criar readiness para Redis e ClamAV, separado do health.", "Disponibilidade"),
    ("L-03", "P0", "Definir e corrigir tratamento de PII ao alternar para anônimo.", "Privacidade"),
    ("L-04", "P0", "Definir significado e regra de sigilo para denúncia identificada.", "Negócio/jurídico"),
    ("L-05", "P0", "Fechar contrato do áudio: captura, anonimização, transcrição, envio e retenção.", "Funcional/privacidade"),
    ("L-06", "P1", "Definir campos obrigatórios por perfil e aplicar validação frontend/backend.", "Negócio"),
    ("L-07", "P1", "Definir idempotência e tratamento de resposta perdida/duplicidade.", "Integração"),
    ("L-08", "P1", "Alinhar tipos, quantidade, tamanho por arquivo e tamanho total dos anexos.", "Operação"),
    ("L-09", "P1", "Homologar Redis e ClamAV reais e executar carga de 100 uploads simultâneos.", "Infraestrutura"),
    ("L-10", "P1", "Definir acompanhamento por protocolo e comprovante.", "Pós-envio"),
    ("L-11", "P1", "Formalizar base legal, aviso de privacidade, retenção e direitos LGPD.", "Jurídico"),
    ("L-12", "P1", "Definir SLO 99,5%, SLIs, alertas e orçamento de erro.", "Operação"),
    ("L-13", "P1", "Adicionar correlation ID, métricas e política de logs por 10 dias.", "Observabilidade"),
    ("L-14", "P1", "Homologar Safari real e executar avaliação manual de acessibilidade/eMAG.", "Qualidade"),
    ("L-15", "P1", "Criar definição de implantação, secrets, probes e rollback.", "DevOps"),
    ("L-16", "P2", "Definir persistência de rascunho e comportamento ao recarregar a página.", "UX"),
    ("L-17", "P2", "Definir timeout/cancelamento no frontend e mensagens específicas por erro.", "Robustez"),
    ("L-18", "P2", "Definir antibot/WAF somente conforme ameaça e política de acesso público.", "Segurança"),
    ("L-19", "P2", "Definir rotina de limpeza de temporários órfãos após crash.", "Operação"),
    ("L-20", "P2", "Corrigir documentação desatualizada e completar OpenAPI.", "Governança"),
]


QUESTION_GROUPS = [
    ("A. Usuários e objetivo do serviço", [
        ("Negócio/Requisitante", "P0", "Quem pode denunciar: cidadão, trabalhador, sindicato, advogado, servidor ou todos?"),
        ("Negócio", "P0", "Quais tipos de manifestação pertencem a este canal e quais devem ser redirecionados?"),
        ("Usuários", "P1", "Quais situações mais comuns motivam uma denúncia e quais informações as pessoas normalmente conhecem?"),
        ("Requisitante", "P1", "A denúncia de órgão público é outro fluxo, outra API ou apenas um link externo?"),
        ("Requisitante", "P1", "Qual destino correto do card da Ouvidoria?"),
        ("Usuários", "P2", "O vídeo explicativo é necessário? Qual conteúdo, duração, Libras e legenda?"),
    ]),
    ("B. Jornada e obrigatoriedade", [
        ("Analista de negócio", "P0", "Quais campos são obrigatórios em cada etapa e sob quais condições?"),
        ("Negócio", "P0", "UF e município sempre são obrigatórios, inclusive para fatos online ou sem local conhecido?"),
        ("Negócio", "P1", "Período, modalidade, número de prejudicados e funções são opcionais ou necessários à triagem?"),
        ("Usuários", "P1", "A navegação deve bloquear avanço ou permitir preenchimento livre e validar apenas no envio?"),
        ("Negócio", "P1", "Ao marcar “Outros”, deve ser obrigatório descrever a irregularidade?"),
        ("Analista de negócio", "P1", "As seis categorias atuais são suficientes e correspondem à classificação interna do MPT?"),
        ("Usuários", "P2", "Deve existir salvamento automático de rascunho? Por quanto tempo e em qual dispositivo?"),
    ]),
    ("C. Áudio e acessibilidade comunicacional", [
        ("Requisitante", "P0", "O relato por áudio é requisito aprovado ou protótipo visual?"),
        ("Negócio/Jurídico", "P0", "O que significa “Sua voz não será identificada”? A voz deve ser alterada ou apenas tratada sob sigilo?"),
        ("Negócio", "P0", "O áudio original deve chegar à API MPT? Em qual formato e limite?"),
        ("Negócio", "P1", "A transcrição é obrigatória? Quem valida erros de transcrição?"),
        ("Jurídico/Privacidade", "P1", "Qual retenção e acesso são permitidos para áudio e transcrição?"),
        ("Usuários", "P2", "Qual alternativa deve existir quando o navegador não permite microfone?"),
    ]),
    ("D. Identificação, anonimato, sigilo e LGPD", [
        ("Jurídico/Negócio", "P0", "Definir precisamente anonimato, sigilo e confidencialidade para este processo."),
        ("Negócio", "P0", "Uma denúncia identificada deve receber sigilo=true ou outro indicador?"),
        ("Privacidade", "P0", "Ao escolher anonimato, todo dado pessoal já digitado deve ser eliminado?"),
        ("Privacidade", "P0", "Qual base legal autoriza o tratamento de dados do denunciante e de terceiros citados?"),
        ("Negócio", "P1", "Quais dados de contato são obrigatórios no modo identificado? Basta um canal de contato?"),
        ("Privacidade", "P1", "Qual aviso de privacidade deve ser mostrado e como registrar ciência quando aplicável?"),
        ("Privacidade", "P1", "Quais prazos de retenção se aplicam a denúncia, anexos, logs e dados de contato?"),
        ("Jurídico", "P1", "Quem pode acessar dados sigilosos e como deve ser auditado esse acesso?"),
        ("Privacidade/Infra", "P1", "IP em logs e rate limit é compatível com a promessa de anonimato?"),
    ]),
    ("E. Empresa e localização", [
        ("Negócio", "P0", "CNPJ precisa ser válido ou pode ser apenas uma referência informada pelo cidadão?"),
        ("Negócio", "P1", "Município deve usar texto livre, lista oficial ou código IBGE?"),
        ("Negócio", "P1", "Como registrar ocorrência sem empresa conhecida, trabalho doméstico ou plataforma digital?"),
        ("Negócio", "P1", "Há roteamento por UF, município, PRT ou unidade interna?"),
        ("Usuários", "P2", "É necessário localizar endereço por CEP ou mapa?"),
    ]),
    ("F. Anexos e segurança", [
        ("Negócio", "P0", "Confirmar tipos de arquivo oficialmente permitidos."),
        ("Negócio/Infra", "P0", "Confirmar 10 arquivos, 20 MB por arquivo e limite total por denúncia."),
        ("Usuários", "P1", "HEIC de iPhone deve ser aceito ou convertido?"),
        ("Segurança", "P0", "Arquivo com MIME/extensão divergente deve ser sempre rejeitado?"),
        ("Negócio", "P1", "Se um anexo estiver infectado, rejeitar toda a denúncia ou permitir remoção e reenvio?"),
        ("Infra", "P1", "Qual tempo máximo aceitável para varrer e transmitir anexos?"),
        ("Segurança", "P1", "É necessário CDR para PDF/DOCX além do antivírus?"),
        ("Operação", "P1", "Como limpar arquivos temporários após crash ou reinício abrupto?"),
    ]),
    ("G. Envio, protocolo e pós-envio", [
        ("Equipe API MPT", "P0", "A API interna suporta Idempotency-Key ou identificador externo estável?"),
        ("Negócio", "P0", "O protocolo gerado pelo BFF é definitivo e reconhecido pelos sistemas internos?"),
        ("Negócio", "P0", "Existe acompanhamento por protocolo? Qual autenticação será exigida?"),
        ("Negócio", "P1", "Deve haver comprovante para download, impressão, e-mail ou SMS?"),
        ("Operação", "P0", "O que fazer se a API MPT aceitar e a resposta se perder?"),
        ("Operação", "P1", "O que fazer em indisponibilidade prolongada: bloquear, enfileirar ou orientar canal alternativo?"),
        ("Usuários", "P1", "Quais mensagens de erro são compreensíveis sem expor detalhes internos?"),
        ("Negócio", "P2", "É necessário permitir complementação posterior da denúncia?"),
    ]),
    ("H. Operação, capacidade e produção", [
        ("Requisitante/Infra", "P0", "Confirmar 100 denúncias/h, 500 usuários, 100 uploads simultâneos e 3 réplicas."),
        ("Requisitante", "P0", "99,5% é SLO aprovado? Qual janela mensal e quais exclusões?"),
        ("Infra", "P0", "Qual plataforma Azure será usada e quem provisionará Redis e ClamAV?"),
        ("Infra", "P0", "Quais sinais tornam uma réplica pronta: Redis, ClamAV e/ou API MPT?"),
        ("Operação", "P1", "Quais thresholds de latência, 5xx, CPU, memória e disco geram alerta?"),
        ("Operação/Privacidade", "P1", "Onde os logs ficarão e como garantir retenção de exatamente 10 dias?"),
        ("Infra", "P1", "Como serão geridos secrets, certificados, DNS, backup e rollback?"),
        ("QA", "P1", "Quais critérios objetivos aprovam o teste de carga e a entrada em produção?"),
    ]),
    ("I. Qualidade, acessibilidade e suporte", [
        ("Usuários/QA", "P0", "Quais perfis de usuário participarão da homologação e quais jornadas serão assinadas?"),
        ("Acessibilidade", "P1", "Além de WCAG 2.1 AA, eMAG é obrigatório?"),
        ("QA", "P1", "Quais versões reais de Safari/iOS e demais navegadores devem ser homologadas?"),
        ("Negócio", "P1", "Há requisito de idiomas, Libras, linguagem simples ou leitura facilitada?"),
        ("Suporte", "P1", "Quem recebe incidentes e dúvidas do canal e em qual horário?"),
        ("Governança", "P1", "Quem é dono de cada requisito e quem aprova mudanças de escopo?"),
    ]),
]


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_status_cells(table, status_column=2):
    colors = {
        "Implementado": "E2F0D9", "Atendido": "E2F0D9",
        "Parcial": "FFF2CC", "Somente visual": "FCE4D6",
        "Ausente": "F4CCCC", "Não atendido": "F4CCCC",
        "Não atendido no código": "F4CCCC", "Não formalizado": "F4CCCC",
        "Não comprovado": "F4CCCC",
    }
    set_repeat_table_header(table.rows[0])
    for row in table.rows[1:]:
        value = row.cells[status_column].text.strip()
        if value in colors:
            set_cell_shading(row.cells[status_column], colors[value])


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("MPT • Canal de Denúncias • Requisitos V2  |  Página ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY_D
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


def add_question_table(doc, title, questions):
    add_heading(doc, title, 2)
    rows = [["Público principal", "Prioridade", "Pergunta"]]
    rows.extend(questions)
    table = add_table(doc, rows, col_widths_cm=[3.2, 1.8, 11.5], header_hex=NAVY_HEX)
    set_repeat_table_header(table.rows[0])
    return table


def main():
    doc = Document()
    doc.core_properties.title = "Relatório Executivo de Requisitos — Canal de Denúncias"
    doc.core_properties.subject = "RF, RNF, regras, conflitos, lacunas e questionário"
    doc.core_properties.author = "Análise automatizada do codebase — revisão humana requerida"
    doc.core_properties.keywords = "MPT, denúncias, requisitos, RF, RNF, LGPD, questionário"

    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    logo = PROJECT / "public" / "logo-mpt.png"
    if logo.exists():
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.add_run().add_picture(str(logo), width=Cm(3.2))

    add_paragraph(doc, "MINISTÉRIO PÚBLICO DO TRABALHO", size=10, bold=True,
                  color=GRAY_D, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    add_paragraph(doc, "Canal de Denúncias da Cidadania", size=11, color=GRAY_D,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
    add_paragraph(doc, "RELATÓRIO EXECUTIVO DE REQUISITOS", size=24, bold=True,
                  color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=3)
    add_paragraph(doc, "Requisitos funcionais e não funcionais • conflitos • lacunas • questionário",
                  size=13, color=TEAL, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
    add_paragraph(doc, "Versão 2.0  |  24/07/2026  |  Base: codebase local e worktree atual",
                  size=10, color=GRAY_D, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    add_paragraph(doc, "Documento para validação com usuários, requisitantes, jurídico/privacidade e analista de negócio",
                  size=10, bold=True, color=GOLD, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=24)

    add_heading(doc, "Controle do documento", 1)
    control = [
        ["Item", "Valor"],
        ["Finalidade", "Consolidar requisitos observados e decisões pendentes antes da produção."],
        ["Escopo analisado", "Frontend Angular, BFF Express, contratos, segurança, testes e documentos locais."],
        ["Não constitui", "Especificação formal aprovada, parecer jurídico ou aceite de produção."],
        ["Fonte de verdade desta versão", "Código local em 24/07/2026; divergências documentais foram registradas."],
        ["Arquivo anterior preservado", "RELATORIO_REQUISITOS_DIRETORIA.docx não foi sobrescrito."],
    ]
    add_table(doc, control, col_widths_cm=[4.3, 12.2])

    add_heading(doc, "Sumário executivo", 1)
    add_paragraph(
        doc,
        "A aplicação implementa o núcleo de registro e encaminhamento de denúncias: wizard, "
        "anonimato/identificação, anexos, revisão, antimalware, integração com a API MPT, "
        "protocolo e tratamento de falhas. A arquitetura é adequada ao porte e está coberta "
        "por testes. Entretanto, o código não equivale a uma especificação de negócio aprovada. "
        "Há promessas visuais e de privacidade que não possuem comportamento equivalente, além "
        "de requisitos operacionais ainda dependentes de infraestrutura e homologação.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )

    rf_counts = Counter(row[2] for row in RF)
    rnf_counts = Counter(row[2] for row in RNF)
    summary = [
        ["Indicador", "Resultado"],
        ["Requisitos funcionais extraídos", str(len(RF))],
        ["RF implementados", str(rf_counts["Implementado"])],
        ["RF parciais", str(rf_counts["Parcial"])],
        ["RF somente visuais ou ausentes", str(rf_counts["Somente visual"] + rf_counts["Ausente"])],
        ["Requisitos não funcionais extraídos", str(len(RNF))],
        ["Conflitos/ambiguidades documentados", str(len(CONFLICTS))],
        ["Lacunas priorizadas", str(len(GAPS))],
        ["Perguntas para stakeholders", str(sum(len(items) for _, items in QUESTION_GROUPS))],
    ]
    add_table(doc, summary, col_widths_cm=[11, 5.5])

    add_heading(doc, "Decisões executivas prioritárias", 2)
    for item in [
        "Definir se áudio é requisito real e o que significa anonimização da voz.",
        "Definir juridicamente e tecnicamente anonimato, sigilo e tratamento de dados pessoais.",
        "Definir campos obrigatórios e impedir avanço/envio inconsistente.",
        "Fechar política de anexos: tipos, limites, assinatura real, infecção e retenção.",
        "Acordar idempotência, protocolo, acompanhamento e indisponibilidade da API MPT.",
        "Formalizar SLO, carga, logs, readiness, infraestrutura Azure e critérios de produção.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Situação para produção", 2)
    add_paragraph(
        doc,
        "Conclusão executiva: apto para homologação controlada, mas não aprovado para produção. "
        "Os bloqueadores são decisões de negócio/privacidade, assinatura real de arquivos, "
        "readiness, idempotência, infraestrutura Redis/ClamAV real, carga, Safari real e aceite "
        "dos usuários responsáveis.",
        size=11, bold=True, color=RED,
    )

    add_heading(doc, "Método e limitações", 1)
    for item in [
        "Leitura do modelo de domínio, componentes, serviços, clientes, middleware, rotas e configuração.",
        "Cruzamento com planejamento.md, spec_design.md, AGENTS.md, README, env.template e relatórios de testes.",
        "Rastreabilidade até arquivos do codebase; requisitos operacionais informados pelo requisitante foram identificados como tal.",
        "Testes reportados em 24/07/2026: 30 backend, 8 frontend e 16 E2E aprovados; não substituem aceite de negócio.",
        "Ausência no código não significa automaticamente que o requisito deve existir: itens marcados como lacuna exigem decisão.",
        "Questões jurídicas e LGPD precisam de validação humana competente.",
    ]:
        add_bullet(doc, item)

    doc.add_page_break()
    add_heading(doc, "1. Visão da solução e atores", 1)
    add_paragraph(doc, "Fluxo atual: cidadão → SPA Angular → BFF Express → ClamAV/Redis → API interna MPT.",
                  size=11, bold=True, color=TEAL)
    actors = [
        ["Ator", "Interesse / responsabilidade"],
        ["Denunciante anônimo", "Relatar fatos sem fornecer identidade; receber protocolo."],
        ["Denunciante identificado", "Relatar fatos, fornecer contato e eventualmente receber retorno."],
        ["Representante/sindicato/advogado", "Público citado na documentação, sem comportamento diferenciado no código."],
        ["Equipe de triagem MPT", "Receber payload e anexos na API interna; classificar e tratar."],
        ["Administrador/Operação", "Manter BFF, Redis, ClamAV, secrets, logs, monitoramento e disponibilidade."],
        ["Analista de negócio", "Fechar regras, obrigatoriedade, taxonomia, mensagens e critérios de aceite."],
        ["Jurídico/Privacidade", "Validar sigilo, anonimato, base legal, transparência, retenção e acesso."],
        ["QA/Acessibilidade", "Validar navegadores, carga, falhas, WCAG/eMAG e homologação."],
    ]
    add_table(doc, actors, col_widths_cm=[4.5, 12])

    add_heading(doc, "2. Requisitos funcionais", 1)
    rf_rows = [["ID", "Requisito", "Situação", "Evidência", "Observação"]]
    rf_rows.extend(RF)
    rf_table = add_table(doc, rf_rows, col_widths_cm=[1.3, 6.1, 2.2, 3.4, 3.5])
    style_status_cells(rf_table)

    doc.add_page_break()
    add_heading(doc, "3. Requisitos não funcionais", 1)
    rnf_rows = [["ID", "Requisito", "Situação", "Evidência", "Observação"]]
    rnf_rows.extend(RNF)
    rnf_table = add_table(doc, rnf_rows, col_widths_cm=[1.4, 6, 2.2, 3.3, 3.6])
    style_status_cells(rnf_table)

    add_heading(doc, "4. Regras de negócio observadas", 1)
    add_paragraph(
        doc,
        "As regras abaixo descrevem o comportamento atual. Elas precisam ser confirmadas; "
        "não devem ser interpretadas automaticamente como regras aprovadas pelo negócio.",
        color=GRAY_D,
    )
    br_rows = [["ID", "Regra observada", "Evidência"]]
    br_rows.extend(BUSINESS_RULES)
    add_table(doc, br_rows, col_widths_cm=[1.5, 10.5, 4.5])

    add_heading(doc, "5. Dicionário de dados da denúncia", 1)
    data_rows = [["Campo", "Tipo", "Obrigatoriedade atual", "Regra/semântica", "Tratamento atual"]]
    data_rows.extend(DATA_FIELDS)
    add_table(doc, data_rows, col_widths_cm=[3.1, 2.2, 3.1, 4.5, 3.6])

    add_heading(doc, "6. Contrato e respostas", 1)
    endpoints = [
        ["Endpoint/status", "Comportamento"],
        ["POST /api/denuncias — 201", "API MPT aceitou; retorna sucesso, protocolo, mensagem e timestamp."],
        ["POST — 400", "JSON inválido, excesso de arquivos ou erro de formato tratado."],
        ["POST — 413", "Arquivo acima de MAX_FILE_SIZE."],
        ["POST — 422", "Campos mínimos ausentes ou anexo infectado."],
        ["POST — 429", "Rate limit excedido."],
        ["POST — 502", "API interna não aceitou/não respondeu corretamente."],
        ["POST — 503", "Configuração MPT ausente ou scanner antimalware indisponível."],
        ["GET /api/denuncias/info", "Informa endpoint, campos, limites e tipos permitidos."],
        ["GET /health", "Liveness simples do processo; não comprova dependências."],
        ["GET /api-docs", "Swagger em desenvolvimento; opt-in explícito em produção."],
    ]
    add_table(doc, endpoints, col_widths_cm=[5.2, 11.3])

    doc.add_page_break()
    add_heading(doc, "7. Requisitos conflitantes e ambiguidades", 1)
    add_paragraph(
        doc,
        "Conflito significa que texto, interface, modelo e comportamento não comunicam a mesma regra. "
        "Os itens críticos devem ser resolvidos antes da homologação com dados reais.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    conflict_rows = [["ID", "Severidade", "Conflito/ambiguidade", "Decisão necessária"]]
    conflict_rows.extend(CONFLICTS)
    conflict_table = add_table(doc, conflict_rows, col_widths_cm=[1.2, 1.8, 7.3, 6.2])
    set_repeat_table_header(conflict_table.rows[0])
    for row in conflict_table.rows[1:]:
        severity = row.cells[1].text.strip()
        set_cell_shading(row.cells[1], "F4CCCC" if severity == "Crítico" else "FFF2CC" if severity == "Alto" else "EDEDED")

    add_heading(doc, "8. Lacunas e backlog de decisão", 1)
    gap_rows = [["ID", "Prioridade", "Lacuna/decisão", "Área"]]
    gap_rows.extend(GAPS)
    gap_table = add_table(doc, gap_rows, col_widths_cm=[1.2, 1.7, 10.2, 3.4])
    set_repeat_table_header(gap_table.rows[0])
    for row in gap_table.rows[1:]:
        priority = row.cells[1].text.strip()
        set_cell_shading(row.cells[1], {"P0": "F4CCCC", "P1": "FFF2CC", "P2": "E2F0D9"}[priority])

    add_heading(doc, "9. Critérios mínimos para produção", 1)
    criteria = [
        ["Dimensão", "Critério de saída"],
        ["Negócio", "P0 respondidos e RF/RN aprovados por dono identificado."],
        ["Privacidade", "Regra de anonimato/sigilo, PII, aviso e retenção formalmente aprovada."],
        ["Upload", "Magic bytes, ClamAV real, política de infecção e temporários validados."],
        ["Integração", "API MPT real homologada, idempotência decidida e falhas testadas."],
        ["Disponibilidade", "/health e /ready separados; 99,5% formalizado com alertas."],
        ["Carga", "100 uploads simultâneos sem OOM/restart, temporário órfão ou bypass de scanner."],
        ["Qualidade", "Suítes aprovadas, Safari real, acessibilidade manual/eMAG conforme decisão."],
        ["Operação", "3 réplicas, Redis/ClamAV privados, secrets, logs 10 dias, rollback e runbook."],
        ["Aceite", "Homologação assinada por negócio, usuários, segurança/privacidade e operação."],
    ]
    add_table(doc, criteria, col_widths_cm=[4, 12.5])

    doc.add_page_break()
    add_heading(doc, "10. Questionário para levantamento e validação", 1)
    add_paragraph(
        doc,
        "Orientação: registrar para cada resposta o decisor, a data, a regra aprovada, o critério "
        "de aceite e o impacto em telas, payload, backend, API MPT e operação. Perguntas P0 são "
        "bloqueadoras; P1 devem ser fechadas antes da produção; P2 podem compor roadmap.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    for title, questions in QUESTION_GROUPS:
        add_question_table(doc, title, questions)

    add_heading(doc, "11. Roteiro sugerido de workshops", 1)
    workshops = [
        ["Sessão", "Participantes", "Saída esperada"],
        ["1 — Escopo e jornada (90 min)", "Usuários, requisitante, negócio", "RF, campos obrigatórios, fluxos alternativos e mensagens."],
        ["2 — Sigilo e privacidade (90 min)", "Negócio, jurídico, privacidade, segurança", "Anonimato, sigilo, PII, áudio, logs e retenção."],
        ["3 — Integração e operação (90 min)", "API MPT, arquitetura, infraestrutura, suporte", "Idempotência, SLO, readiness, Redis, ClamAV e incidentes."],
        ["4 — Aceite e priorização (60 min)", "Todos os donos", "Backlog P0/P1/P2, critérios de homologação e responsáveis."],
    ]
    add_table(doc, workshops, col_widths_cm=[4.2, 5.2, 7.1])

    add_heading(doc, "12. Recomendação executiva", 1)
    add_paragraph(
        doc,
        "Suspender novas refatorações estruturais e investir na validação dos requisitos. "
        "O melhor retorno agora vem de resolver os P0, corrigir promessas conflitantes, homologar "
        "dependências reais e obter aceite formal. Após os workshops, converter as respostas em "
        "histórias/regras com critérios de aceite e atualizar modelo TypeScript, serialização, "
        "validação backend, OpenAPI, testes e documentação de forma rastreável.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )

    add_heading(doc, "Referências analisadas", 2)
    for ref in [
        "planejamento.md — objetivo, público e jornada desejada.",
        "spec_design.md — design system, responsividade e padrões de UI.",
        "AGENTS.md — arquitetura, guardrails e pontos desatualizados registrados como conflitos.",
        "src/app/models/complaint.model.ts — campos e enumerações.",
        "src/app/components/** — comportamento das etapas e revisão.",
        "src/app/services e clients — estado, submissão e protocolo.",
        "server/controllers, services, clients, middleware, routes, infrastructure e index.js.",
        "server/env.template, server/README.md e Swagger.",
        "e2e/**, specs frontend/backend e RELATORIO_TESTES_DENUNCIAS.md.",
        "RELATORIO_FINAL_SANITIZACAO_CODEBASE_DENUNCIAS.md.",
    ]:
        add_bullet(doc, ref)

    add_paragraph(
        doc,
        "Fim do relatório — versão 2.0. Recomenda-se revisão humana antes de distribuição externa.",
        size=9, bold=True, color=GRAY_D, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=16,
    )

    doc.save(OUT)
    print(f"Gerado: {OUT}")
    print(f"RF: {len(RF)} | RNF: {len(RNF)} | Conflitos: {len(CONFLICTS)} | Perguntas: {sum(len(q) for _, q in QUESTION_GROUPS)}")


if __name__ == "__main__":
    main()
