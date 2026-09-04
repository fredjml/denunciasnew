"""Gera RELATORIO_EXECUTIVO_DIRETORIA.docx com aparência executiva."""
from pathlib import Path
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).parent / "RELATORIO_EXECUTIVO_DIRETORIA.docx"

# Paleta institucional
NAVY = RGBColor(0x17, 0x33, 0x5C)
TEAL = RGBColor(0x0C, 0x6E, 0x64)
GOLD = RGBColor(0xBF, 0x95, 0x3F)
GRAY_D = RGBColor(0x5A, 0x5A, 0x5A)
GRAY_L_HEX = "F0F0F0"
NAVY_HEX = "17335C"
TEAL_HEX = "0C6E64"
GREEN = RGBColor(0x2E, 0x7D, 0x32)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_borders(cell, color: str = "17335C", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)
        tc_borders.append(border)
    tc_pr.append(tc_borders)


def add_paragraph(doc, text, *, size=11, bold=False, color=None,
                  align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.alignment = align
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 20, 2: 15, 3: 12}
    colors = {1: NAVY, 2: TEAL, 3: NAVY}
    p = add_paragraph(doc, text, size=sizes[level], bold=True,
                      color=colors[level], space_before=12, space_after=6)
    # Underline for level 1
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), NAVY_HEX)
        pBdr.append(bottom)
        pPr.append(pBdr)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    return p


def add_table(doc, rows, col_widths_cm=None):
    n_rows = len(rows)
    n_cols = len(rows[0])
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for r_idx, row_data in enumerate(rows):
        for c_idx, value in enumerate(row_data):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            # Clear default paragraph and add ours
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            run = p.add_run(str(value))
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            if r_idx == 0:
                set_cell_shading(cell, NAVY_HEX)
                run.font.color.rgb = WHITE
                run.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                if r_idx % 2 == 0:
                    set_cell_shading(cell, GRAY_L_HEX)
            set_cell_borders(cell)

    if col_widths_cm:
        for c_idx, width in enumerate(col_widths_cm):
            for row in table.rows:
                row.cells[c_idx].width = Cm(width)
    return table


def main():
    doc = Document()

    # Margens
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    # Fonte padrão
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ======== CABEÇALHO ========
    add_paragraph(doc, "MINISTÉRIO PÚBLICO DO TRABALHO",
                  size=10, bold=True, color=GRAY_D,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    add_paragraph(doc, "Canal de Denúncias da Cidadania",
                  size=10, color=GRAY_D,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    add_paragraph(doc, "Relatório Executivo",
                  size=26, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    add_paragraph(doc, "Sanitização e endurecimento do codebase",
                  size=14, color=TEAL,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    add_paragraph(doc,
                  "Julho de 2026  |  Ambiente: repositório local  |  Branch: refactor/frontend-architecture",
                  size=10, color=GRAY_D,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    # ======== SUMÁRIO EXECUTIVO ========
    add_heading(doc, "Sumário executivo", 1)
    add_paragraph(doc,
        "O Canal de Denúncias passou por uma rodada de sanitização proporcional, focada em "
        "eliminar falhas de correção, reduzir superfície de ataque e preparar o serviço para os parâmetros "
        "de produção informados (100 denúncias/hora, 500 usuários, 100 uploads simultâneos, 3 réplicas e "
        "Safari obrigatório). O trabalho foi executado localmente, sem alterar o desenho da interface "
        "aprovado pelos usuários responsáveis pelo negócio.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_paragraph(doc,
        "Foram alterados 56 arquivos (+6.544 / −1.308 linhas), organizados em 17 incrementos "
        "priorizados por custo-benefício.",
        bold=True, color=NAVY)

    # ======== INDICADORES ========
    add_heading(doc, "Indicadores mensuráveis", 2)
    indicadores = [
        ["Indicador", "Antes", "Depois"],
        ["Bundle inicial do frontend", "≈ 601 kB", "384 kB  (−36%)"],
        ["Transferência estimada", "acima do orçamento", "85 kB  (limite 120 kB)"],
        ["Vulnerabilidades em produção (frontend)", "> 0", "0"],
        ["Vulnerabilidades em produção (backend)", "> 0", "0"],
        ["Testes automatizados backend", "praticamente inexistentes", "30 / 30"],
        ["Testes unitários frontend", "limitados", "8 / 8"],
        ["Jornadas E2E (Playwright + axe)", "0", "16 / 16"],
        ["Navegadores validados", "somente Chromium", "Chromium, Firefox e WebKit/Safari"],
        ["Rate limit com 3 réplicas", "3 contadores independentes", "contador único via Redis"],
        ["Verificação antimalware nos uploads", "não havia", "obrigatória (ClamAV / INSTREAM)"],
    ]
    add_table(doc, indicadores, col_widths_cm=[7, 4, 6])

    # ======== ENTREGAS ========
    add_heading(doc, "Principais entregas por eixo", 2)

    add_heading(doc, "Correção do fluxo de denúncia", 3)
    add_bullet(doc, "Eliminado o falso sucesso: o protocolo passa a ser exibido somente após confirmação real da API do MPT.")
    add_bullet(doc, "Contrato HTTP alinhado: 201 apenas em aceite; 502 para falha da integração; 503 quando dependência não configurada.")
    add_bullet(doc, "A tela de revisão mantém os dados preenchidos e oferece nova tentativa em caso de falha.")

    add_heading(doc, "Segurança e privacidade", 3)
    add_bullet(doc, "Upload agora é feito em disco temporário e verificado por antivírus antes de sair do servidor.")
    add_bullet(doc, "CORS restrito ao domínio do frontend, Swagger desligado em produção e inicialização abortada quando faltam variáveis essenciais.")
    add_bullet(doc, "Logs deixaram de expor UF, sigilo e mensagens brutas da integração; URLs e logs de áudio foram removidos do frontend.")
    add_bullet(doc, "Zero vulnerabilidades conhecidas em dependências de produção, no frontend e no backend.")

    add_heading(doc, "Preparação para operar em 3 réplicas", 3)
    add_bullet(doc, "Rate limit compartilhado entre réplicas via Redis (redis@6.1.0 + rate-limit-redis@6.0.0).")
    add_bullet(doc, "Falha no Redis bloqueia submissões e encerra o servidor, evitando comportamento inconsistente entre nós.")

    add_heading(doc, "Qualidade e cobertura de testes", 3)
    add_bullet(doc, "ESLint no frontend e no backend com zero erros.")
    add_bullet(doc, "Vitest e Supertest cobrindo validação, protocolo, controller, clientes HTTP, perímetro e integração de upload seguro.")
    add_bullet(doc, "Playwright + axe automatizando a jornada completa em Chromium (desktop e mobile), Firefox e WebKit (Safari).")

    add_heading(doc, "Desempenho e experiência", 3)
    add_bullet(doc, "Substituição do Bootstrap completo por módulos Sass necessários: bundle inicial reduzido em 36%.")
    add_bullet(doc, "Ajustes de acessibilidade (teclado, foco, ARIA, contraste e semântica) sem alterar UI/UX aprovada.")

    # ======== PENDÊNCIAS ========
    add_heading(doc, "Pendências para operar em produção", 2)
    add_paragraph(doc,
        "O código está apto a seguir para homologação. A conclusão até produção depende de "
        "dependências operacionais externas ao repositório:")
    pendencias = [
        ["Item", "Responsável sugerido", "Bloqueia produção?"],
        ["Provisionar Azure Managed Redis (ou equivalente) e testar REDIS_URL com 3 réplicas", "Infraestrutura", "Sim"],
        ["Provisionar ClamAV em rede privada com StreamMaxLength > 20 MB", "Infraestrutura / Segurança", "Sim"],
        ["Testar API interna do MPT em homologação (timeout, indisponibilidade, idempotência)", "Integrações", "Sim"],
        ["Teste de carga em staging com 100 uploads simultâneos", "Qualidade / Infra", "Sim"],
        ["Homologação em Safari real (macOS / iPhone)", "Qualidade", "Sim"],
        ["Retenção de logs de 10 dias e revisão de exportação de dados sensíveis", "Segurança", "Sim"],
        ["Plataforma de hospedagem, certificados, DNS, secrets e monitoramento", "DevOps", "Sim"],
        ["Homologação funcional pelos responsáveis pelo negócio", "Negócio", "Sim"],
    ]
    add_table(doc, pendencias, col_widths_cm=[9, 4.5, 3.5])

    # ======== RISCOS ACEITOS ========
    add_heading(doc, "Riscos aceitos e não bloqueantes", 2)
    add_bullet(doc, "Depreciação de @import no Sass usado pelo Bootstrap 5.3 — o build passa e a migração não traz retorno proporcional agora.")
    add_bullet(doc, "Avisos moderados na árvore de desenvolvimento do frontend (Angular/CLI) — dependências de produção sem vulnerabilidades.")
    add_bullet(doc, "Avisos documentais em OpenAPI (licença, servidor localhost, operationId, resposta 4xx).")

    # ======== RECOMENDAÇÃO ========
    add_heading(doc, "Recomendação da equipe técnica", 2)
    add_paragraph(doc,
        "A sanitização proporcional foi concluída e não recomendamos novas refatorações estruturais "
        "neste momento. O próximo investimento com melhor retorno é homologar infraestrutura, integração "
        "e carga em ambiente real, e não continuar reorganizando o código.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    add_paragraph(doc, "Status técnico: Pronto para homologação.",
                  size=12, bold=True, color=GREEN, space_before=6)
    add_paragraph(doc, "Status para produção: Aguardando dependências externas listadas acima.",
                  size=12, bold=True, color=GOLD)

    # ======== REFERÊNCIAS ========
    add_heading(doc, "Documentos de referência", 3)
    add_bullet(doc, "Analises/RELATORIO_FINAL_SANITIZACAO_CODEBASE_DENUNCIAS.md — detalhamento dos 17 incrementos.")
    add_bullet(doc, "Analises/RELATORIO_ARQUIVOS_ALTERADOS_DENUNCIAS.md — melhoria arquivo por arquivo.")

    doc.save(OUT)
    print(f"Gerado: {OUT}")


if __name__ == "__main__":
    main()
