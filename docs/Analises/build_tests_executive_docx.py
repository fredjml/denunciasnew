"""Gera RELATORIO_EXECUTIVO_TESTES_DIRETORIA.docx a partir do RELATORIO_TESTES_DENUNCIAS.md."""
from pathlib import Path
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).parent / "RELATORIO_EXECUTIVO_TESTES_DIRETORIA.docx"

NAVY = RGBColor(0x17, 0x33, 0x5C)
TEAL = RGBColor(0x0C, 0x6E, 0x64)
GOLD = RGBColor(0xBF, 0x95, 0x3F)
GRAY_D = RGBColor(0x5A, 0x5A, 0x5A)
GRAY_L_HEX = "F0F0F0"
NAVY_HEX = "17335C"
GREEN = RGBColor(0x2E, 0x7D, 0x32)
GREEN_HEX = "2E7D32"
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_borders(cell, color: str = "17335C", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)
        tc_borders.append(border)
    tc_pr.append(tc_borders)


def add_paragraph(doc, text, *, size=11, bold=False, color=None,
                  align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.alignment = align
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 20, 2: 15, 3: 12}
    colors = {1: NAVY, 2: TEAL, 3: NAVY}
    p = add_paragraph(doc, text, size=sizes[level], bold=True,
                      color=colors[level], space_before=12, space_after=6)
    if level == 1:
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), NAVY_HEX)
        pBdr.append(bottom)
        pPr.append(pBdr)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    return p


def add_table(doc, rows, col_widths_cm=None, header_hex=NAVY_HEX):
    n_rows = len(rows)
    n_cols = len(rows[0])
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for r_idx, row_data in enumerate(rows):
        for c_idx, value in enumerate(row_data):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            run = p.add_run(str(value))
            run.font.name = "Calibri"
            run.font.size = Pt(10)
            if r_idx == 0:
                set_cell_shading(cell, header_hex)
                run.font.color.rgb = WHITE
                run.bold = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                if r_idx % 2 == 0:
                    set_cell_shading(cell, GRAY_L_HEX)
                # Cor verde para colunas com "0" (falhas) ou marcadores ✓
                text = str(value).strip()
                if text in ("✅", "✓", "0"):
                    run.font.color.rgb = GREEN
                    run.bold = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_borders(cell)

    if col_widths_cm:
        for c_idx, width in enumerate(col_widths_cm):
            for row in table.rows:
                row.cells[c_idx].width = Cm(width)
    return table


def main():
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # ======== CAPA ========
    add_paragraph(doc, "MINISTÉRIO PÚBLICO DO TRABALHO",
                  size=10, bold=True, color=GRAY_D,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    add_paragraph(doc, "Canal de Denúncias da Cidadania",
                  size=10, color=GRAY_D,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    add_paragraph(doc, "Relatório Executivo de Testes",
                  size=26, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    add_paragraph(doc, "Unitários, integração, E2E e regressão",
                  size=14, color=TEAL,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    add_paragraph(doc,
                  "Execução local em 24/07/2026  |  Node 22.22.3  |  Branch: refactor/frontend-architecture",
                  size=10, color=GRAY_D,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)

    # ======== SUMÁRIO ========
    add_heading(doc, "Sumário executivo", 1)
    add_paragraph(doc,
        "Todas as suítes de teste foram executadas na branch de refatoração e passaram "
        "sem falhas. A régua de testes cobre correção do fluxo (falso sucesso), contrato HTTP, "
        "upload seguro com antivírus, perímetro do backend, arquitetura em camadas, "
        "acessibilidade WCAG e regressão em três motores de navegador.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    add_paragraph(doc, "Resultado consolidado: 54 casos executados, 54 aprovados, 0 falhas.",
                  bold=True, color=GREEN, size=12, space_before=6)

    add_heading(doc, "Painel de resultados", 2)
    painel = [
        ["Suíte", "Ferramenta", "Total", "OK", "Falhas"],
        ["Unitários — Backend",        "Vitest 4.1.10",             "22", "22", "0"],
        ["Integração — Backend",       "Vitest + Supertest",        "8",  "8",  "0"],
        ["Unitários — Frontend",       "Vitest via Angular",        "8",  "8",  "0"],
        ["E2E funcional + regressão",  "Playwright",                "12", "12", "0"],
        ["E2E acessibilidade",         "Playwright + axe-core",     "4",  "4",  "0"],
        ["Lint — Frontend",            "ESLint",                    "—",  "0 erros", "0"],
        ["Lint — Backend",             "ESLint",                    "—",  "0 erros", "0"],
        ["Auditoria produção — Frontend", "npm audit --omit=dev",   "99 deps", "0 vulns", "0"],
        ["Auditoria produção — Backend",  "npm audit --omit=dev",   "150 deps","0 vulns", "0"],
    ]
    add_table(doc, painel, col_widths_cm=[5, 4.5, 2.5, 2.5, 2])

    # ======== BACKEND ========
    add_heading(doc, "Backend — Vitest (30 casos, 1,79 s)", 1)

    add_heading(doc, "Regras da denúncia (services)", 3)
    for t in [
        "Gera protocolo com formato esperado.",
        "Permite geração determinística de protocolo em testes.",
        "Aceita denúncia completa.",
        "Reporta ausência de localização e descrição obrigatória.",
        "Aceita relato escrito quando nenhuma irregularidade é marcada.",
        "Constrói payload de integração e deriva sigilo/anonimato.",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "Cliente ClamAV (antimalware)", 3)
    for t in [
        "Encaminha arquivo limpo usando o framing oficial INSTREAM.",
        "Rejeita arquivo quando o ClamAV reporta ameaça.",
        "Falha fechada quando o scanner não está configurado.",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "Cliente da API interna do MPT", 3)
    for t in [
        "Rejeita requisição quando MPT_API_URL não está configurada.",
        "Envia denúncia como multipart com timeout configurado.",
        "Encaminha autenticação e anexos para a API MPT.",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "Store de rate limit em Redis (3 réplicas)", 3)
    for t in [
        "Usa store em memória fora de produção quando Redis não é configurado.",
        "Conecta store externo usando URL e prefixo configurados.",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "Middleware de upload seguro", 3)
    add_bullet(doc, "Grava anexos fora da memória e remove arquivo temporário após uso.")

    add_heading(doc, "Controller de recebimento (integração)", 3)
    for t in [
        "400 quando o campo denúncia não é JSON válido.",
        "422 quando dados obrigatórios estão faltando.",
        "201 apenas quando a API interna do MPT aceita.",
        "502 sem resposta de sucesso quando a API interna falha.",
        "503 sem resposta de sucesso quando MPT_API_URL está ausente.",
        "Rejeita anexo quando o antivírus reporta ameaça.",
        "Falha fechada quando a varredura de anexos não está disponível.",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "Upload seguro — integração ponta a ponta", 3)
    add_bullet(doc, "Endpoint Express real → disco temporário → ClamAV TCP simulado → multipart → API MPT HTTP simulada.")

    add_heading(doc, "Perímetro do servidor (index)", 3)
    for t in [
        "Health endpoint fica fora do rate limit de submissão.",
        "Limita submissões repetidas do endpoint de denúncia.",
        "Aceita apenas a origem configurada em produção (CORS).",
        "Swagger fica desabilitado por padrão em produção.",
        "Permite opt-in explícito para o Swagger em produção.",
        "Não sobe em produção sem configuração obrigatória.",
        "Sobe em produção usando o store de rate limit compartilhado.",
    ]:
        add_bullet(doc, t)

    # ======== FRONTEND ========
    add_heading(doc, "Frontend — Vitest via Angular (8 casos, 15,89 s)", 1)

    add_heading(doc, "Cliente HTTP de denúncia", 3)
    for t in [
        "Serializa dados e anexos como multipart.",
        "Rejeita resposta HTTP sem sucesso.",
        "Rejeita resposta de sucesso sem protocolo válido (evita falso sucesso).",
    ]:
        add_bullet(doc, t)

    add_heading(doc, "Serviço de estado da denúncia", 3)
    for t in [
        "Inicia no passo de boas-vindas com valores padrão.",
        "Atualiza dados e mantém navegação entre passos dentro dos limites.",
        "Envia a denúncia atual e guarda o protocolo devolvido.",
        "Expõe falha visível após erro do cliente — sem simular sucesso.",
        "Limpa a falha visível quando uma nova tentativa é bem-sucedida.",
    ]:
        add_bullet(doc, t)

    # ======== E2E ========
    add_heading(doc, "E2E — Playwright + axe-core (16 execuções)", 1)

    add_paragraph(doc,
        "Última corrida oficial: test-results/.last-run.json → status: passed, sem falhas.",
        color=GRAY_D, size=10)

    matriz = [
        ["Cenário", "Tipo", "Chrome Desktop", "Chrome Mobile", "Firefox", "Safari (WebKit)"],
        ["Preenche, navega e conclui denúncia simulada", "Funcional", "✅", "✅", "✅", "✅"],
        ["Mantém revisão, exibe falha e permite nova tentativa", "Regressão", "✅", "✅", "✅", "✅"],
        ["Não confirma denúncia sem protocolo válido", "Regressão", "✅", "✅", "✅", "✅"],
        ["Baseline automatizada de acessibilidade", "Acessibilidade", "✅", "✅", "✅", "✅"],
    ]
    add_table(doc, matriz, col_widths_cm=[5.5, 2.5, 2, 2, 2, 2.5])

    add_heading(doc, "Cobertura observada nos E2E", 3)
    for t in [
        "Preenchimento, navegação, upload de PDF e leitura do protocolo devolvido.",
        "Ao receber 503, alerta é exibido e revisão preserva os dados; nova tentativa com 201 conclui.",
        "201 sem protocolo válido não confirma a denúncia; alerta é exibido.",
        "Análise axe com regras WCAG 2.0/2.1 A e AA — nenhuma violação critical/serious.",
    ]:
        add_bullet(doc, t)

    # ======== REGRESSÃO ========
    add_heading(doc, "Testes de regressão — o que cada bloco protege", 1)
    regress = [
        ["Incremento", "Regressão evitada", "Testes que protegem"],
        ["3", "Falso sucesso no frontend", "Frontend cliente/serviço + E2E de falha e ausência de protocolo"],
        ["4", "Contrato HTTP 201/422/502/503", "Backend controller + E2E funcional e sem protocolo"],
        ["11", "Regras extraídas do controller", "Backend services + controller aprovando 201"],
        ["12", "Perímetro (rate limit, CORS, Swagger, fail-fast)", "Backend index (7 casos)"],
        ["13", "Cliente HTTP no frontend", "Frontend cliente (3 casos)"],
        ["16", "Rate limit consistente entre 3 réplicas", "Store Redis + perímetro em produção"],
        ["17", "Upload seguro com antimalware", "ClamAV + middleware + controller (ameaça/indisponível) + integração"],
        ["6",  "Acessibilidade WCAG", "E2E axe (4 execuções)"],
        ["8",  "Jornada real da denúncia", "E2E funcional (4 execuções)"],
        ["14/15", "Cobertura Firefox e WebKit/Safari", "Projetos firefox-desktop e webkit-desktop"],
    ]
    add_table(doc, regress, col_widths_cm=[2, 6, 8])

    # ======== QUALIDADE ESTÁTICA ========
    add_heading(doc, "Qualidade estática e dependências", 1)
    qs = [
        ["Verificação", "Comando", "Resultado"],
        ["Lint frontend",    "eslint src/**/*.{ts,html}",  "0 erros / 0 avisos"],
        ["Lint backend",     "eslint **/*.js (server)",    "0 erros / 0 avisos"],
        ["Vulnerabilidades frontend (produção)", "npm audit --omit=dev", "0 em 99 deps"],
        ["Vulnerabilidades backend (produção)",  "npm audit --omit=dev", "0 em 150 deps"],
        ["Contrato OpenAPI", "redocly lint",               "válido; apenas avisos documentais"],
    ]
    add_table(doc, qs, col_widths_cm=[6, 6, 4])

    # ======== CONCLUSÃO ========
    add_heading(doc, "Conclusão", 1)
    add_paragraph(doc,
        "Não há falhas de teste conhecidas. A cobertura automatizada existe para os cenários "
        "críticos de correção, contrato e segurança do fluxo de denúncia, e é executada "
        "em três motores de navegador antes de qualquer entrega.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    add_paragraph(doc, "Status técnico: 54/54 aprovados.",
                  size=12, bold=True, color=GREEN, space_before=6)
    add_paragraph(doc,
        "Pendências não cobertas por testes automatizados dependem de infraestrutura real "
        "(Redis, ClamAV, API interna do MPT, Safari em macOS/iOS, teste de carga em staging).",
        size=11, color=GOLD, bold=True)

    add_heading(doc, "Documentos de referência", 3)
    for t in [
        "Analises/RELATORIO_TESTES_DENUNCIAS.md — detalhamento por caso.",
        "Analises/vitest-backend.json — saída bruta da suíte backend.",
        "cidadania-canal-denuncias/test-results/ — vídeos, screenshots e anexos por projeto Playwright.",
        "cidadania-canal-denuncias/playwright-report/index.html — relatório navegável.",
    ]:
        add_bullet(doc, t)

    doc.save(OUT)
    print(f"Gerado: {OUT}")


if __name__ == "__main__":
    main()
